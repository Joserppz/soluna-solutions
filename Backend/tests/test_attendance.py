"""Asistencia de los sábados (diseño págs. 3 y 10): marcar, cerrar, personas nuevas y resumen."""

import csv
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core import clock
from app.models import AuditLog, Inquiry
from tests.conftest import requires_db

pytestmark = requires_db

SCHEDULES = "/api/v1/schedules"


def activity(client, headers, days=0, publish=True, title="Movimiento en compañía"):
    response = client.post(
        SCHEDULES,
        json={
            "title": title,
            "category": "fisioterapia",
            "date": (clock.today() + timedelta(days=days)).isoformat(),
            "startTime": "10:00",
            "endTime": "11:30",
            "location": "Sede ASPAR",
            "publish": publish,
        },
        headers=headers,
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def url(schedule_id, suffix=""):
    return f"{SCHEDULES}/{schedule_id}/attendance{suffix}"


@pytest.fixture
def today_id(client, auth_headers):
    return activity(client, auth_headers("coordination"))


def mark(client, headers, schedule_id, patient_id, status, note=None):
    body = {"status": status} | ({"note": note} if note is not None else {})
    return client.put(url(schedule_id, f"/{patient_id}"), json=body, headers=headers)


class TestTomarAsistencia:
    def test_la_lista_son_los_afiliados_activos_por_nombre(self, client, auth_headers, today_id):
        sheet = client.get(url(today_id), headers=auth_headers("volunteer")).json()

        assert [item["name"] for item in sheet["items"]] == [
            "Elena Rojas",
            "Gloria Condori",
            "Miguel Fernández",
        ]
        assert sheet["summary"] == {"came": 0, "notCame": 0, "unmarked": 3, "total": 3}
        assert sheet["closed"] is False
        assert sheet["activity"]["title"] == "Movimiento en compañía"
        # La lista no tiene datos clínicos
        assert set(sheet["items"][0]) == {"patientId", "name", "status", "note", "markedAt"}

    def test_marcar_vino_y_no_vino_con_nota(self, client, auth_headers, today_id, patient_ids):
        volunteer = auth_headers("volunteer")

        came = mark(
            client, volunteer, today_id, patient_ids["consent"], "vino", "Viene con su hijo"
        )
        absent = mark(client, volunteer, today_id, patient_ids["no_consent"], "no_vino")
        again = mark(client, volunteer, today_id, patient_ids["consent"], "no_vino")

        assert came.status_code == 200, came.text
        assert came.json()["note"] == "Viene con su hijo" and came.json()["markedAt"]
        assert absent.json()["status"] == "no_vino"
        # Sin "note" se conserva la nota anterior
        assert again.json()["note"] == "Viene con su hijo"
        summary = client.get(url(today_id), headers=volunteer).json()["summary"]
        assert summary == {"came": 0, "notCame": 2, "unmarked": 1, "total": 3}

    def test_buscar_por_nombre(self, client, auth_headers, today_id):
        sheet = client.get(
            url(today_id), params={"search": "GLORIA"}, headers=auth_headers("volunteer")
        ).json()

        assert [item["name"] for item in sheet["items"]] == ["Gloria Condori"]

    def test_marcar_a_todos_no_pisa_a_quien_no_vino(
        self, client, auth_headers, today_id, patient_ids
    ):
        volunteer = auth_headers("volunteer")
        mark(client, volunteer, today_id, patient_ids["no_consent"], "no_vino")

        sheet = client.post(url(today_id, "/mark-all-present"), headers=volunteer).json()

        statuses = {item["name"]: item["status"] for item in sheet["items"]}
        assert statuses == {
            "Elena Rojas": "vino",
            "Gloria Condori": "vino",
            "Miguel Fernández": "no_vino",
        }
        assert sheet["summary"]["came"] == 2

    def test_cerrar_la_asistencia_y_reabrirla(self, client, auth_headers, today_id, patient_ids):
        volunteer = auth_headers("volunteer")
        mark(client, volunteer, today_id, patient_ids["consent"], "vino")

        closed = client.post(url(today_id, "/close"), headers=volunteer)
        late = mark(client, volunteer, today_id, patient_ids["no_consent"], "vino")
        twice = client.post(url(today_id, "/close"), headers=volunteer)

        assert closed.json()["closed"] is True and closed.json()["closedAt"]
        assert late.status_code == 409 and "ya se cerró" in late.json()["detail"]
        assert twice.status_code == 409
        assert client.post(url(today_id, "/reopen"), headers=volunteer).status_code == 403
        reopened = client.post(url(today_id, "/reopen"), headers=auth_headers("coordination"))
        assert reopened.json()["closed"] is False
        assert (
            mark(client, volunteer, today_id, patient_ids["no_consent"], "vino").status_code == 200
        )

    def test_cerrar_queda_en_la_bitacora(self, client, auth_headers, today_id, db_session):
        client.post(url(today_id, "/close"), headers=auth_headers("volunteer"))

        log = db_session.scalar(
            select(AuditLog).where(AuditLog.action == "Cierre de asistencia del día")
        )
        assert log.category == "ASISTENCIA"

    @pytest.mark.parametrize(
        ("days", "publish", "cancel", "message"),
        [
            (0, False, False, "borrador"),
            (0, True, True, "cancelada"),
            (3, True, False, "el día de la actividad"),
        ],
    )
    def test_solo_actividades_publicadas_cuya_fecha_llego(
        self, client, auth_headers, patient_ids, days, publish, cancel, message
    ):
        coordination = auth_headers("coordination")
        schedule_id = activity(client, coordination, days=days, publish=publish)
        if cancel:
            client.post(
                f"{SCHEDULES}/{schedule_id}/cancel", json={"note": "Lluvia"}, headers=coordination
            )

        response = mark(
            client, auth_headers("volunteer"), schedule_id, patient_ids["consent"], "vino"
        )

        assert response.status_code == 409
        assert message in response.json()["detail"]

    def test_fichas_dadas_de_baja(self, client, auth_headers, today_id, patient_ids):
        volunteer = auth_headers("volunteer")
        mark(client, volunteer, today_id, patient_ids["consent"], "vino")
        client.post(
            f"/api/v1/patients/{patient_ids['consent']}/deactivate",
            json={"reason": "Se mudó de ciudad"},
            headers=auth_headers(),
        )
        client.post(
            f"/api/v1/patients/{patient_ids['revoked']}/deactivate",
            json={"reason": "Se mudó de ciudad"},
            headers=auth_headers(),
        )

        names = [i["name"] for i in client.get(url(today_id), headers=volunteer).json()["items"]]
        response = mark(client, volunteer, today_id, patient_ids["revoked"], "vino")

        # Quien ya tenía asistencia sigue en la lista; quien no, sale
        assert names == ["Elena Rojas", "Miguel Fernández"]
        assert response.status_code == 409

    def test_permisos(self, client, auth_headers, today_id, patient_ids):
        assert client.get(url(today_id), headers=auth_headers(None)).status_code == 403
        assert (
            mark(client, auth_headers(None), today_id, patient_ids["consent"], "vino").status_code
            == 403
        )
        assert client.get(url(today_id), headers=auth_headers("reception")).status_code == 200

    def test_actividad_o_afiliado_inexistente(self, client, auth_headers, today_id):
        import uuid

        volunteer = auth_headers("volunteer")
        assert client.get(url(uuid.uuid4()), headers=volunteer).status_code == 404
        assert mark(client, volunteer, today_id, uuid.uuid4(), "vino").status_code == 404
        bad = client.put(
            url(today_id, f"/{uuid.uuid4()}"), json={"status": "tarde"}, headers=volunteer
        )
        assert bad.status_code == 422


class TestActividadConAsistencia:
    def test_la_agenda_muestra_cuantos_vinieron(self, client, auth_headers, today_id, patient_ids):
        volunteer = auth_headers("volunteer")
        mark(client, volunteer, today_id, patient_ids["consent"], "vino")
        mark(client, volunteer, today_id, patient_ids["no_consent"], "vino")

        detail = client.get(f"{SCHEDULES}/{today_id}", headers=volunteer).json()

        assert detail["came"] == 2

    def test_con_asistencia_no_se_cancela_ni_vuelve_a_borrador(
        self, client, auth_headers, today_id, patient_ids
    ):
        coordination = auth_headers("coordination")
        mark(client, auth_headers("volunteer"), today_id, patient_ids["consent"], "vino")

        cancel = client.post(
            f"{SCHEDULES}/{today_id}/cancel", json={"note": "Lluvia"}, headers=coordination
        )
        unpublish = client.post(f"{SCHEDULES}/{today_id}/unpublish", headers=coordination)

        assert cancel.status_code == 409 and "asistencia registrada" in cancel.json()["detail"]
        assert unpublish.status_code == 409

    def test_un_borrador_con_filas_de_asistencia_no_se_elimina(
        self, client, auth_headers, today_id, patient_ids
    ):
        coordination = auth_headers("coordination")
        mark(client, auth_headers("volunteer"), today_id, patient_ids["consent"], "sin_marcar")
        assert (
            client.post(f"{SCHEDULES}/{today_id}/unpublish", headers=coordination).status_code
            == 200
        )

        response = client.delete(f"{SCHEDULES}/{today_id}", headers=coordination)

        assert response.status_code == 409


class TestPersonaNueva:
    def test_anotar_crea_una_solicitud_para_recepcion(
        self, client, auth_headers, today_id, db_session
    ):
        response = client.post(
            url(today_id, "/newcomers"),
            json={
                "name": "Elena Quisbert",
                "phone": "71234567",
                "onBehalfOf": "Su esposo, Mario",
                "note": "Primera vez · llegó con su esposo",
                "contactConsent": True,
            },
            headers=auth_headers("volunteer"),
        )

        assert response.status_code == 201, response.text
        assert "Recepción abrirá su ficha" in response.json()["message"]
        inquiry = db_session.get(Inquiry, response.json()["inquiryId"])
        assert (inquiry.kind, inquiry.source, inquiry.status) == (
            "asistencia",
            "asistencia",
            "recibida",
        )
        assert str(inquiry.schedule_id) == today_id
        assert inquiry.on_behalf_of == "Su esposo, Mario"
        assert inquiry.contact_consent is True
        assert db_session.scalar(
            select(AuditLog).where(AuditLog.action == "Persona nueva anotada en asistencia")
        )

    def test_tambien_despues_de_cerrar(self, client, auth_headers, today_id):
        volunteer = auth_headers("volunteer")
        client.post(url(today_id, "/close"), headers=volunteer)

        response = client.post(
            url(today_id, "/newcomers"),
            json={"name": "Hilda Pacheco", "phone": "70000011"},
            headers=volunteer,
        )

        assert response.status_code == 201

    @pytest.mark.parametrize(("name", "phone"), [("Al", "70000011"), ("Hilda Pacheco", "123")])
    def test_validaciones(self, client, auth_headers, today_id, name, phone):
        response = client.post(
            url(today_id, "/newcomers"),
            json={"name": name, "phone": phone},
            headers=auth_headers("volunteer"),
        )

        assert response.status_code == 422


class TestResumenMensual:
    @pytest.fixture
    def month(self, client, auth_headers, today_id, patient_ids):
        volunteer = auth_headers("volunteer")
        coordination = auth_headers("coordination")
        mark(client, volunteer, today_id, patient_ids["consent"], "vino")
        mark(client, volunteer, today_id, patient_ids["no_consent"], "vino")
        mark(client, volunteer, today_id, patient_ids["revoked"], "no_vino")
        art = activity(client, coordination, title='=HYPERLINK("http://x")')
        mark(client, volunteer, art, patient_ids["consent"], "vino")
        return clock.today().strftime("%Y-%m")

    def summary(self, client, headers, month, file_format="json"):
        return client.get(
            "/api/v1/attendance/monthly",
            params={"month": month, "format": file_format},
            headers=headers,
        )

    def test_resumen_por_actividad_y_por_tipo(self, client, auth_headers, month, today_id):
        body = self.summary(client, auth_headers("volunteer"), month).json()

        mine = next(a for a in body["activities"] if a["scheduleId"] == today_id)
        assert (mine["came"], mine["notCame"], mine["closed"]) == (2, 1, False)
        physio = next(c for c in body["byCategory"] if c["category"] == "fisioterapia")
        assert physio["label"] == "Fisioterapia" and physio["came"] >= 3
        assert body["totalCame"] == sum(c["came"] for c in body["byCategory"])
        assert body["month"] == month

    def test_csv_para_excel_en_espanol(self, client, auth_headers, month):
        response = self.summary(client, auth_headers("volunteer"), month, "csv")

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert f'filename="asistencia-{month}.csv"' in response.headers["content-disposition"]
        assert response.content.startswith("﻿".encode())
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
        assert rows[0] == [
            "Fecha",
            "Actividad",
            "Tipo",
            "Vinieron",
            "No vinieron",
            "Asistencia cerrada",
        ]
        titles = [row[1] for row in rows[1:] if len(row) == 6]
        # Un título que empieza con "=" no se ejecuta como fórmula al abrirlo
        assert '\'=HYPERLINK("http://x")' in titles
        assert ["Total de asistencias", str(rows[-1][1])] == rows[-1]

    def test_excel(self, client, auth_headers, month):
        response = self.summary(client, auth_headers("coordination"), month, "xlsx")

        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet.title == f"Asistencia {month}"
        assert [cell.value for cell in sheet[1]] == [
            "Fecha",
            "Actividad",
            "Tipo",
            "Vinieron",
            "No vinieron",
            "Asistencia cerrada",
        ]
        assert sheet.cell(row=1, column=1).font.bold

    def test_mes_por_defecto_invalido_y_permisos(self, client, auth_headers):
        assert (
            client.get("/api/v1/attendance/monthly", headers=auth_headers("volunteer")).status_code
            == 200
        )
        assert self.summary(client, auth_headers(), "2026-13").status_code == 422
        assert self.summary(client, auth_headers(), "octubre").status_code == 422
        assert self.summary(client, auth_headers(None), "2026-10").status_code == 403

    def test_un_mes_sin_actividades(self, client, auth_headers):
        body = self.summary(client, auth_headers(), "2020-02").json()

        assert body == {"month": "2020-02", "activities": [], "byCategory": [], "totalCame": 0}


def test_crear_filas_dos_veces_no_choca_con_la_restriccion_unica(db_session, today_id, patient_ids):
    """Dos voluntarios marcando a la misma persona a la vez: el segundo no recibe un 500."""
    from app.models import Attendance
    from app.modules.attendance import repository

    for _ in range(2):
        repository.ensure_rows(
            db_session, today_id, [patient_ids["consent"], patient_ids["revoked"]]
        )

    rows = db_session.scalars(select(Attendance).where(Attendance.schedule_id == today_id)).all()
    assert len(rows) == 2
    assert {row.status for row in rows} == {"sin_marcar"}
